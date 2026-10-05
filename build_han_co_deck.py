from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

ROOT = Path(r"C:\Users\Admin\Desktop\P Kinh")
OUT = ROOT / "outputs"
ASSETS = OUT / "han_co_2_5d_assets"
OUT.mkdir(exist_ok=True)
ASSETS.mkdir(exist_ok=True)

SHEETS = [
    Path(r"C:\Users\Admin\.codex\generated_images\01a0dcd2-b187-76f1-b209-2476fefb86a3\exec-c73fce59-ad4e-4ca5-be73-71d8b46a922b.png"),
    Path(r"C:\Users\Admin\.codex\generated_images\01a0dcd2-b187-76f1-b209-2476fefb86a3\exec-54fbaf5a-ebff-4a27-8419-6a9a22bd0237.png"),
    Path(r"C:\Users\Admin\.codex\generated_images\01a0dcd2-b187-76f1-b209-2476fefb86a3\exec-95b08b87-fb37-4a9c-93f2-7b9e4b495851.png"),
    Path(r"C:\Users\Admin\.codex\generated_images\01a0dcd2-b187-76f1-b209-2476fefb86a3\exec-70e4d9e9-df6d-4185-97ba-5f5ec09018de.png"),
    Path(r"C:\Users\Admin\.codex\generated_images\01a0dcd2-b187-76f1-b209-2476fefb86a3\exec-d5816a65-ff84-4c75-84f1-5250589ed9b6.png"),
]

DATA = [
    ("渴思飲飢思食。渴時飲茶，飢時食飯。", "Khát tư ẩm, cơ tư thực; khát thời ẩm trà, cơ thời thực phạn", "Khát nghĩ đến uống, đói nghĩ đến ăn; khi khát thì uống trà, khi đói thì ăn cơm.", "Đói nhớ ăn, khát nhớ uống: đáp ứng đúng nhu cầu."),
    ("兩岸間，架板橋。橋上行人，橋下行船。", "Lưỡng ngạn gian, giá bản kiều. Kiều thượng hành nhân, kiều hạ hành thuyền", "Giữa hai bờ sông, bắc cây cầu ván. Trên cầu người đi, dưới cầu thuyền đi.", "Nhớ 2 tầng: trên cầu là người, dưới cầu là thuyền."),
    ("庭前樹有鳥巢。小鳥一群，樹間飛鳴。", "Đình tiền thụ, hữu điểu sào. Tiểu điểu nhất quần, thụ gian phi minh", "Trên cây trước sân có tổ chim. Một bầy chim nhỏ bay và kêu giữa các cây.", "Sân – cây – tổ chim – đàn chim."),
    ("畫一幅，馬八匹。或臥，或立，或俯或仰。", "Họa nhất bức, mã bát thất. Hoặc ngọa, hoặc lập, hoặc phủ, hoặc ngưỡng", "Một bức họa có tám con ngựa: con nằm, con đứng, con cúi, con ngước.", "8 ngựa, 4 tư thế: nằm – đứng – cúi – ngước."),
    ("左右手，共十指。左五指，右五指。", "Tả hữu thủ, cộng thập chỉ. Tả ngũ chỉ, hữu ngũ chỉ", "Tay trái, tay phải gồm mười ngón. Mỗi tay có năm ngón.", "Hai bàn tay = 10 ngón; mỗi tay = 5 ngón."),
    ("左右手，能取物，能作事。", "Tả hữu thủ, năng thủ vật, năng tác sự", "Tay trái, tay phải có thể lấy đồ vật, có thể làm công việc.", "Hai công năng: lấy vật và làm việc."),
    ("人之身體有三部分。頭軀幹與四肢。", "Nhân chi thân thể hữu tam bộ phận. Đầu, khu cán dữ tứ chi", "Thân thể người có ba phần: đầu, mình và chân tay.", "Đầu – mình – chân tay."),
    ("頭在上，軀幹居中，兩足在下，兩手在兩旁。", "Đầu tại thượng, khu cán cư trung, lưỡng túc tại hạ, lưỡng thủ tại lưỡng bàng", "Đầu ở trên, mình ở giữa, hai chân ở dưới, hai tay ở hai bên.", "Nhớ vị trí từ trên xuống và hai bên."),
    ("人面上部為顙，極下為頷，鼻居中央。", "Nhân diện thượng bộ vi tảng, cực hạ vi hạm, tị cư trung ương", "Trán ở phần trên, cằm ở phần dưới cùng, mũi ở chính giữa.", "Trán trên – mũi giữa – cằm dưới."),
    ("鼻下有口，口中有舌，鼻上有兩目。", "Tị hạ hữu khẩu, khẩu trung hữu thiệt, tị thượng hữu lưỡng mục", "Dưới mũi có miệng, trong miệng có lưỡi; trên mũi có hai con mắt.", "Mắt trên – mũi giữa – miệng dưới."),
    ("目上有眉，兩耳在面之左右邊。", "Mục thượng hữu mi; lưỡng nhĩ tại diện chi tả hữu biên", "Trên mắt có lông mày; hai tai ở bên trái và phải của mặt.", "Mắt + lông mày; tai ở hai bên."),
    ("腦主傳令；手足承腦之命令以行動。", "Não chủ truyền lệnh; thủ túc thừa não chi mệnh lệnh dĩ hành động", "Óc truyền lệnh; tay chân nhận lệnh rồi hành động.", "Não ra lệnh, tay chân thực hiện."),
    ("腦能接認感覺，亦能發生思想。", "Não năng tiếp nhận cảm giác, diệc năng phát sinh tư tưởng", "Óc tiếp nhận cảm giác và phát sinh tư tưởng.", "Não nhận cảm giác và tạo suy nghĩ."),
    ("心司發血，肺主呼吸。", "Tâm ti phát huyết, phế chủ hô hấp", "Tim truyền máu, phổi chủ về hô hấp.", "Tim → máu; phổi → hô hấp."),
    ("肝，脾，胃，小腸均屬消化之機官。", "Can, tì, vị, tiểu trường quân thuộc tiêu hóa chi cơ quan", "Gan, tì, bao tử và ruột non đều thuộc cơ quan tiêu hóa.", "Nhóm tiêu hóa: gan – tì – vị – ruột non."),
    ("大腸，腎與膀胱皆任排泄之役。", "Đại trường, thận dữ bàng quang giai nhậm bài tiết chi dịch", "Ruột già, thận và bàng quang đảm nhiệm việc bài tiết.", "Nhóm bài tiết: ruột già – thận – bàng quang."),
    ("願以此功德，普及於一切，我等與眾生，皆共成佛道。", "Nguyện dĩ thử công đức, phổ cập ư nhất thiết, ngã đẳng dữ chúng sanh, giai cộng thành Phật đạo", "Nguyện đem công đức này hướng về khắp tất cả, để mọi chúng sanh cùng thành Phật đạo.", "Công đức hướng đến tất cả chúng sanh."),
    ("觀自在菩薩，行深般若波羅蜜多時，照見五蘊皆空，度一切苦厄。", "Quán Tự Tại Bồ Tát hành thâm Bát nhã Ba la mật đa thời, chiếu kiến ngũ uẩn giai không, độ nhất thiết khổ ách", "Quán Tự Tại quán chiếu sâu, thấy năm uẩn là không và vượt qua khổ ách.", "Quán chiếu sâu → thấy không → vượt khổ."),
    ("世間一切萬物，皆有定理。", "Thế gian nhất thiết vạn vật, giai hữu định lý", "Tất cả muôn vật trong thế gian đều có lý nhất định.", "Muôn vật đều có quy luật."),
    ("世間一切萬物，皆有軌範。", "Thế gian nhất thiết vạn vật, giai hữu quỹ phạm", "Tất cả muôn vật đều có khuôn mẫu của nó.", "Mỗi sự vật có khuôn phép riêng."),
    ("人類亦然，其法尤周密。", "Nhân loại diệc nhiên, kỳ pháp vưu chu mật", "Con người cũng vậy, và quy tắc xã hội còn chặt chẽ hơn.", "Con người sống theo quy tắc."),
    ("故各國聖人所說之法。", "Cố các quốc thánh nhân, sở thuyết chi pháp", "Vì vậy, các bậc Thánh nhân ở các nước đều nói ra những pháp phù hợp cho con người.", "Bậc hiền triết chỉ dạy con đường sống."),
    ("皆有使人行入正軌之法。", "Giai hữu sử nhân hành nhập chánh quỹ chi pháp", "Đều có phương pháp giúp con người đi vào khuôn phép đúng đắn.", "Đi vào đường lối đúng đắn."),
    ("如禮教，法律，規約等是也。", "Như lễ giáo, pháp luật, quy ước đẳng thị dã", "Như lễ giáo, pháp luật, quy ước, v.v.", "Lễ giáo – pháp luật – quy ước."),
    ("我佛世尊所說之法。", "Ngã Phật Thế tôn sở thuyết chi pháp", "Pháp do Đức Phật Thế Tôn tuyên nói.", "Nhấn mạnh giáo pháp của Đức Phật."),
    ("然於人生法，亦大致相同。", "Nhiên ư nhân sanh pháp, diệc đại trí tương đồng", "Với phương pháp dạy người đời, đại ý cũng có những điểm tương đồng.", "Giáo pháp và phép dạy đời có điểm tương đồng."),
    ("其尤要者，教吾人守五戒。", "Kỳ vưu yếu giả, giáo ngô nhân thủ ngũ giới", "Điều cốt yếu là dạy chúng ta giữ gìn năm giới.", "Cốt yếu: giữ Năm giới."),
    ("行八正道，以不失人格。", "Hành bát chánh đạo, dĩ bất thất nhân cách", "Thực hành Bát Chánh Đạo để giữ gìn nhân cách làm người.", "Bát Chánh Đạo giúp giữ nhân cách."),
    ("故吾人應當學佛法。", "Cố ngô nhân ưng đương học Phật pháp", "Vì những lý do trên, chúng ta nên học Phật pháp.", "Kết luận: nên học Phật pháp."),
]

PALETTES = [
    ("F8F1E5", "1C3854", "E0A83A"), ("F7F3EB", "4A2B53", "D48D32"), ("EEF5F5", "20535A", "F19C79"),
    ("FBF4E6", "6A3A20", "D5A24B"), ("F2F0FA", "3A3158", "B889D4"), ("F1F6EC", "35563A", "9DCA7A"),
]

def rgb(hexv):
    return RGBColor(int(hexv[0:2],16), int(hexv[2:4],16), int(hexv[4:6],16))

def crop_panels():
    result = []
    for sheet_i, sheet_path in enumerate(SHEETS):
        img = Image.open(sheet_path).convert('RGB')
        w, h = img.size
        for r in range(2):
            for c in range(3):
                idx = sheet_i * 6 + r * 3 + c
                l, t = int(c*w/3)+5, int(r*h/2)+5
                rr, bb = int((c+1)*w/3)-5, int((r+1)*h/2)-5
                crop = img.crop((l,t,rr,bb))
                out = ASSETS / f"scene-{idx+1:02d}.jpg"
                crop.save(out, quality=95)
                result.append(out)
    return result

def add_textbox(slide, x, y, w, h, text, size, color, bold=False, align=PP_ALIGN.LEFT, font='Aptos'):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear(); tf.word_wrap = True; tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = align
    run = p.add_run(); run.text = text; run.font.name = font; run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = color
    return box

def add_panel(slide, x, y, w, h, fill, radius=True):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape

def add_picture_cover(slide, path, x,y,w,h):
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))

def slide_bg(slide, color):
    bg = slide.background.fill
    bg.solid(); bg.fore_color.rgb = color

def make_deck(scene_paths):
    prs = Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    blank = prs.slide_layouts[6]
    # Cover
    s = prs.slides.add_slide(blank); slide_bg(s, rgb('1B2D3A'))
    add_picture_cover(s, scene_paths[29], 6.9, 0, 6.43, 7.5)
    add_panel(s, 0.55, 0.6, 5.8, 5.7, rgb('1B2D3A'), False)
    add_textbox(s, 0.75, 1.0, 5.2, 0.8, 'HÁN CỔ', 34, rgb('E7BC55'), True)
    add_textbox(s, 0.75, 1.8, 5.4, 1.6, '29 CÂU\nDỄ NHỚ', 36, rgb('FFFFFF'), True)
    add_textbox(s, 0.78, 4.05, 5.1, 0.9, 'Học bằng hình ảnh 2.5D\nHán tự · Âm Hán Việt · Nghĩa · Gợi nhớ', 17, rgb('D7E6E8'))
    add_textbox(s, 0.78, 6.65, 5.0, 0.35, 'Ôn tập cuối học kỳ 2 · Khóa IX · Khoa PHTX', 10, rgb('B2C4C8'))

    for i, (han, hv, meaning, mn) in enumerate(DATA):
        s = prs.slides.add_slide(blank)
        bg_hex, dark_hex, accent_hex = PALETTES[i % len(PALETTES)]
        bgc, dark, accent = rgb(bg_hex), rgb(dark_hex), rgb(accent_hex)
        slide_bg(s, bgc)
        n = i + 1; layout = i % 6; img = scene_paths[i]
        # unique top number and section tag
        add_textbox(s, 0.55, 0.28, 2.5, 0.35, f'CÂU {n:02d}', 13, accent, True)
        add_textbox(s, 10.9, 0.28, 1.8, 0.35, 'HÁN CỔ', 10, dark, True, PP_ALIGN.RIGHT)
        if layout == 0:
            add_panel(s, 0.5, 0.85, 5.2, 5.9, rgb('FFFFFF'))
            add_picture_cover(s, img, 7.1, 0.7, 5.7, 6.1)
            tx,ty,tw = 0.85,1.18,4.5
        elif layout == 1:
            add_picture_cover(s, img, 0.5, 0.85, 5.2, 5.65)
            add_panel(s, 6.05, 1.0, 6.75, 5.8, rgb('FFFFFF'))
            tx,ty,tw = 6.45,1.35,5.95
        elif layout == 2:
            add_panel(s, 0.5, 0.85, 12.3, 1.45, rgb('FFFFFF'))
            add_picture_cover(s, img, 0.55, 2.65, 5.5, 4.25)
            add_panel(s, 6.45, 2.65, 6.35, 4.25, rgb('FFFFFF'))
            tx,ty,tw = 6.85,2.98,5.55
        elif layout == 3:
            add_panel(s, 0.48, 0.85, 12.35, 5.9, rgb('FFFFFF'))
            add_picture_cover(s, img, 0.55, 0.92, 4.25, 5.75)
            tx,ty,tw = 5.25,1.2,7.0
        elif layout == 4:
            add_picture_cover(s, img, 7.5, 1.1, 5.15, 5.45)
            add_panel(s, 0.55, 0.95, 6.45, 5.75, rgb('FFFFFF'))
            tx,ty,tw = 0.95,1.28,5.65
        else:
            add_panel(s, 0.5, 0.85, 4.45, 5.95, rgb('FFFFFF'))
            add_picture_cover(s, img, 5.35, 0.85, 7.45, 3.15)
            add_panel(s, 5.35, 4.35, 7.45, 2.45, rgb('FFFFFF'))
            tx,ty,tw = 0.9,1.2,3.65
        # Text content
        add_textbox(s, tx, ty, tw, 0.75, han, 19, dark, True, font='Microsoft JhengHei')
        add_textbox(s, tx, ty+0.88, tw, 0.7, 'Hán Việt: ' + hv, 11.5, dark)
        add_textbox(s, tx, ty+1.78, tw, 0.9, 'Nghĩa: ' + meaning, 13, dark)
        note = add_panel(s, tx, ty+3.0, tw, 1.05, rgb('FFF7E0'))
        add_textbox(s, tx+0.18, ty+3.18, tw-0.36, 0.6, 'GỢI NHỚ  ' + mn, 13, accent, True)
        add_textbox(s, 0.55, 7.08, 12.1, 0.2, f'{n:02d} / 29    •    Hán tự  |  Âm Hán Việt  |  Nghĩa  |  Gợi nhớ', 8.5, dark, False, PP_ALIGN.CENTER)
    output = OUT / 'Han_co_29_cau_de_nho_2_5D.pptx'
    prs.save(output)
    return output

if __name__ == '__main__':
    scenes = crop_panels()
    result = make_deck(scenes)
    print(result)
